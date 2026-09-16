from pathlib import Path
from io import BytesIO
from zipfile import ZipFile

from langchain_community.document_loaders import (
    PyPDFLoader,
    Docx2txtLoader,
    TextLoader
)

from pptx import Presentation
from PIL import Image
from pypdf import PdfReader

from ocr import perform_ocr


# Folder containing all documents
DOCUMENT_FOLDER = Path("../documents/TechNova/")


def load_txt(file_path):
    """
    Load a TXT file using LangChain.
    """

    loader = TextLoader(
        str(file_path),
        encoding="utf-8"
    )

    documents = loader.load()

    return [
        {
            "content": documents[0].page_content,
            "metadata": {
                "file_name": file_path.name,
                "file_type": file_path.suffix,
                "source": str(file_path)
            }
        }
    ]


def load_pdf(file_path):
    """
    Load PDF text and OCR embedded images.

    Normal PDF text is extracted using PyPDFLoader.
    Images embedded inside PDF pages are sent to OCR.
    """

    # Extract normal PDF text
    loader = PyPDFLoader(str(file_path))

    documents = loader.load()

    # Open PDF to access embedded images
    reader = PdfReader(str(file_path))

    pdf_documents = []

    for page_number, document in enumerate(
        documents,
        start=1
    ):

        # Native PDF text
        text = document.page_content.strip()

        # Get corresponding PDF page
        pdf_page = reader.pages[page_number - 1]

        ocr_texts = []

        # Check for embedded images
        for image in pdf_page.images:

            try:

                image_data = image.data

                pil_image = Image.open(
                    BytesIO(image_data)
                )

                # OCR the embedded image
                image_text = perform_ocr(
                    pil_image
                )

                if image_text:
                    ocr_texts.append(
                        image_text
                    )

            except Exception as e:

                print(
                    f"OCR failed for image "
                    f"on page {page_number}: {e}"
                )

        # Combine native text + OCR text
        if ocr_texts:

            text += "\n\n" + "\n\n".join(
                ocr_texts
            )

        pdf_documents.append(
            {
                "content": text,

                "metadata": {
                    "file_name": file_path.name,
                    "file_type": file_path.suffix,
                    "page": page_number,
                    "source": str(file_path)
                }
            }
        )

    return pdf_documents


def load_docx(file_path):
    """
    Load DOCX native text and OCR embedded images.
    """

    loader = Docx2txtLoader(str(file_path))

    documents = loader.load()

    content = documents[0].page_content.strip()

    # DOCX files are ZIP containers
    with ZipFile(file_path, "r") as zip_file:

        image_files = [
            name
            for name in zip_file.namelist()
            if name.startswith("word/media/")
        ]

        for image_file in image_files:

            try:

                image_data = zip_file.read(
                    image_file
                )

                image = Image.open(
                    BytesIO(image_data)
                )

                # OCR embedded image
                ocr_text = perform_ocr(image)

                if ocr_text:
                    content += (
                        "\n\n" + ocr_text
                    )

            except Exception as e:

                print(
                    f"OCR failed for image "
                    f"{image_file}: {e}"
                )

    return [
        {
            "content": content,

            "metadata": {
                "file_name": file_path.name,
                "file_type": file_path.suffix,
                "source": str(file_path)
            }
        }
    ]


def load_pptx(file_path):
    """
    Load PowerPoint slide text and OCR images.
    """

    presentation = Presentation(file_path)

    ppt_documents = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1
    ):

        slide_content = []

        for shape in slide.shapes:

            # Normal slide text
            if hasattr(shape, "text"):

                text = shape.text.strip()

                if text:
                    slide_content.append(
                        text
                    )

            # Embedded image
            if shape.shape_type == 13:

                try:

                    image = Image.open(
                        BytesIO(
                            shape.image.blob
                        )
                    )

                    # OCR image
                    ocr_text = perform_ocr(
                        image
                    )

                    if ocr_text:
                        slide_content.append(
                            ocr_text
                        )

                except Exception as e:

                    print(
                        f"OCR failed for image "
                        f"on slide {slide_number}: {e}"
                    )

        ppt_documents.append(
            {
                "content": "\n\n".join(
                    slide_content
                ),

                "metadata": {
                    "file_name": file_path.name,
                    "file_type": file_path.suffix,
                    "slide": slide_number,
                    "source": str(file_path)
                }
            }
        )

    return ppt_documents


def load_image(file_path):
    """
    Load an image and extract text using OCR.
    """

    image = Image.open(file_path)

    text = perform_ocr(image)

    return [
        {
            "content": text,

            "metadata": {
                "file_name": file_path.name,
                "file_type": file_path.suffix,
                "source": str(file_path)
            }
        }
    ]


def load_documents():
    """
    Scan the documents folder and load all supported files.
    """

    all_documents = []

    for file_path in DOCUMENT_FOLDER.iterdir():

        if not file_path.is_file():
            continue

        extension = file_path.suffix.lower()

        if extension == ".pdf":

            all_documents.extend(
                load_pdf(file_path)
            )

        elif extension == ".docx":

            all_documents.extend(
                load_docx(file_path)
            )

        elif extension == ".txt":

            all_documents.extend(
                load_txt(file_path)
            )

        elif extension == ".pptx":

            all_documents.extend(
                load_pptx(file_path)
            )

        elif extension in [
            ".png",
            ".jpg",
            ".jpeg"
        ]:

            all_documents.extend(
                load_image(file_path)
            )

    return all_documents